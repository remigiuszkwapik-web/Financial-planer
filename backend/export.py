"""Export the ledger to Excel (.xlsx) or Markdown (.md).

Both formats are generated from the same deterministic summary, so they always
agree with each other and with the app.
"""

from __future__ import annotations

import io

from .amount import format_cents
from .calc import summarize


def to_markdown(transactions: list[dict]) -> str:
    summary = summarize(transactions)
    lines: list[str] = []
    lines.append("# Finanzübersicht\n")

    lines.append("## Summen pro Kategorie\n")
    lines.append("| Kategorie | Anzahl | Summe |")
    lines.append("| --- | ---: | ---: |")
    for cat in summary["categories"]:
        lines.append(
            f"| {cat['name']} | {cat['count']} | {format_cents(cat['total_cents'])} |"
        )
    u = summary["unsorted"]
    if u["count"]:
        lines.append(f"| _(unsortiert)_ | {u['count']} | {format_cents(u['total_cents'])} |")
    lines.append(
        f"| **Gesamt** | | **{format_cents(summary['grand_total_cents'])}** |\n"
    )

    lines.append("## Alle Buchungen\n")
    lines.append("| Kategorie | Empfänger | Betrag | Datum |")
    lines.append("| --- | --- | ---: | --- |")
    for tx in transactions:
        cat = tx.get("category_name") or "—"
        payee = (tx.get("payee") or "").replace("|", "/")
        date = (tx.get("created_at") or "")[:10]
        lines.append(f"| {cat} | {payee} | {format_cents(tx['amount_cents'])} | {date} |")

    return "\n".join(lines) + "\n"


def to_xlsx(transactions: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    summary = summarize(transactions)
    wb = Workbook()

    # Sheet 1: summary
    ws = wb.active
    ws.title = "Übersicht"
    ws.append(["Kategorie", "Anzahl", "Summe (€)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for cat in summary["categories"]:
        ws.append([cat["name"], cat["count"], cat["total_cents"] / 100])
    u = summary["unsorted"]
    if u["count"]:
        ws.append(["(unsortiert)", u["count"], u["total_cents"] / 100])
    total_row = ["Gesamt", "", summary["grand_total_cents"] / 100]
    ws.append(total_row)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)

    # Sheet 2: all transactions
    ws2 = wb.create_sheet("Buchungen")
    ws2.append(["Kategorie", "Empfänger", "Betrag (€)", "Datum"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for tx in transactions:
        ws2.append([
            tx.get("category_name") or "",
            tx.get("payee") or "",
            tx["amount_cents"] / 100,
            (tx.get("created_at") or "")[:10],
        ])

    for sheet in (ws, ws2):
        for col in sheet.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = max(12, width + 2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
